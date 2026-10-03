#!/usr/bin/env python3
"""Builds eval/results/skbouncer-eval-1.1.xlsx: data sheets from the raw results
and a Criteria sheet whose figures are formulas over them.

usage: eval/workbook.py RESULTS OUT.xlsx   (needs openpyxl; RESULTS as for eval/score.py)"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "eval"))
import score  # noqa: E402

R = Path(sys.argv[1])
OUT = Path(sys.argv[2])

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEAD_FILL = PatternFill("solid", start_color="D9E1F2")
wb = Workbook()
wb.remove(wb.active)


def hexes(v):
    return " ".join(f"{x:04x}" for x in v)


def sheet(name, header, rows, widths=None):
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append(list(r))
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(40, len(h) + 2)))
    col = {h: get_column_letter(i) for i, h in enumerate(header, 1)}
    return ws, col, len(rows) + 1


def rng(name, col, last, h):
    return f"{name}!${col[h]}$2:${col[h]}${last}"


# -- Exp1 ---------------------------------------------------------------------
e1 = json.loads((R / "exp1/results.json").read_text())
ws1, c1, n1 = sheet("Exp1", ["No", "Function", "Evidence", "Result"],
                    [(r["no"], r["function"], r["evidence"], "PASS" if r["pass"] else "FAIL") for r in e1["rows"]],
                    {"Function": 34, "Evidence": 80})
c3 = e1["c3"]
ws1["F1"], ws1["G1"] = "C3 traffic", "Value"
for i, (k, v) in enumerate([("SSH connections", c3["ssh_connections"]), ("Random-byte connections", c3["random_connections"]),
                            ("Connections received", c3["connections_received"]), ("Bytes received", c3["bytes_received"]),
                            ("skbouncer events", c3["events"])], 2):
    ws1[f"F{i}"], ws1[f"G{i}"] = k, v
ws1.column_dimensions["F"].width = 26

# -- Exp3 ---------------------------------------------------------------------
h3 = ["Run", "Client", "Container", "Port", "Server record_len", "Server legacy_version", "Server ciphers",
      "Server versions", "Events", "Event container", "Event record_len", "Event legacy_version", "captured_len",
      "ciphers_seen", "versions_seen", "Event ciphers", "Event versions", "Verdict", "Reference verdict",
      "record_len match", "legacy match", "Cipher match", "Version match", "Named correctly", "Verdict agrees",
      "Grade (skbouncer)", "Grade (server)", "Read in full", "Grade agrees"]
rows3 = []
for part, pol, c in score.exp3_connections(R):
    s, e = c["server"], c["events"][0]
    rows3.append([part, c["label"], c["container"], c["port"], s["record_len"], s["legacy_version"], hexes(s["ciphers"]),
                  hexes(s["versions"]), len(c["events"]), e.get("container"), e["record_len"], e["legacy_version"],
                  e["captured_len"], e["ciphers_seen"], e["versions_seen"], hexes(e["ciphers"]), hexes(e["versions"]),
                  e["verdict"], score.reference_verdict(pol, c["container"], s), None, None, None, None, None, None,
                  score.grade(e["ciphers"], e["versions"], e["legacy_version"]),
                  score.grade(s["ciphers"], s["versions"], s["legacy_version"]), None, None])
ws3, c3c, n3 = sheet("Exp3", h3, rows3, {"Client": 34, "Container": 26, "Server ciphers": 50, "Event ciphers": 50})
for i in range(2, n3 + 1):
    L = lambda h: f"{c3c[h]}{i}"
    ws3[L("record_len match")] = f"={L('Event record_len')}={L('Server record_len')}"
    ws3[L("legacy match")] = f"={L('Event legacy_version')}={L('Server legacy_version')}"
    ws3[L("Cipher match")] = f'=IF({L("ciphers_seen")},EXACT({L("Event ciphers")},{L("Server ciphers")}),"")'
    ws3[L("Version match")] = f'=IF({L("versions_seen")},EXACT({L("Event versions")},{L("Server versions")}),"")'
    ws3[L("Named correctly")] = f"={L('Event container')}={L('Container')}"
    ws3[L("Verdict agrees")] = f'=IF({L("Verdict")}="unknown","",{L("Verdict")}={L("Reference verdict")})'
    ws3[L("Read in full")] = f'=AND({L("ciphers_seen")},{L("versions_seen")},{L("Run")}<>"census-mixed")'
    ws3[L("Grade agrees")] = f'=IF({L("Read in full")},{L("Grade (skbouncer)")}={L("Grade (server)")},"")'
ws3["S1"].comment = Comment("Computed by eval/score.py reference_verdict(): policy/README.md applied to the "
                            "server's own decode of the ClientHello.", "eval")
ws3[f"{c3c['Grade (server)']}1"].comment = Comment("Durumeric et al., NDSS 2017, section V-A; rules in the "
                                                    "eval/score.py docstring.", "eval")

# -- Exp3 fuzz ----------------------------------------------------------------
fz = score.jsonl(R / "exp3/fuzz/mutants.jsonl")
rowsf = []
for r in fz:
    truth = score.decode_mutant(bytes.fromhex(r["hex"]))
    wrong = 0
    for e in r["events"]:
        for f, seen in (("ciphers", "ciphers_seen"), ("versions", "versions_seen")):
            if e.get(seen) and truth[f] is not None and e.get(f) and e.get(f) != truth[f]:
                wrong += 1
    rowsf.append([r["i"], r["kind"], r["from"], len(r["hex"]) // 2, r["received"], len(r["events"]),
                  sum(bool(e.get("ciphers_seen")) for e in r["events"]), sum(bool(e.get("versions_seen")) for e in r["events"]),
                  wrong])
wsf, cf, nf = sheet("Exp3_fuzz", ["i", "Kind", "From", "Bytes sent", "Bytes received", "Events",
                                  "Events with ciphers_seen", "Events with versions_seen", "False complete fields"],
                    rowsf, {"From": 34})
frun = json.loads((R / "exp3/fuzz/run.json").read_text())
for i, (k, v) in enumerate([("skbouncer alive at end", frun["alive_at_end"]), ("skb_tc attached at end", frun["attached_at_end"]),
                            ("Detected", frun["counters"]["detected"]), ("Too short to classify", frun["counters"]["short"])], 2):
    wsf[f"K{i}"], wsf[f"L{i}"] = k, v
wsf["K1"], wsf["L1"] = "Run", "Value"
wsf.column_dimensions["K"].width = 26

# -- Exp4 ---------------------------------------------------------------------
pq = score.post_quantum()
rows4 = []
for arm in sorted(p.name for p in (R / "exp4").iterdir() if p.is_dir()):
    run = json.loads((R / f"exp4/{arm}/run.json").read_text())
    for c in score.jsonl(R / f"exp4/{arm}/connections.jsonl"):
        e = c["events"][0]
        rows4.append([arm, run["mtu"], run["timestamps"], run["mss"], c["label"], pq.get(c["label"], False), c["rep"],
                      c["server"]["record_len"], e["record_len"], e["captured_len"], None, None,
                      e["ciphers_seen"], e["versions_seen"], e["verdict"], None])
h4 = ["Arm", "MTU", "Timestamps", "MSS", "Client", "Post-quantum", "Rep", "Server record_len", "Event record_len",
      "captured_len", "Split (truth)", "Flagged split", "ciphers_seen", "versions_seen", "Verdict", "Decidable"]
ws4, c4, n4 = sheet("Exp4", h4, rows4, {"Client": 34})
for i in range(2, n4 + 1):
    L = lambda h: f"{c4[h]}{i}"
    ws4[L("Split (truth)")] = f"={L('Server record_len')}+5>{L('MSS')}"
    ws4[L("Flagged split")] = f"={L('captured_len')}<{L('Event record_len')}"
    ws4[L("Decidable")] = f'={L("Verdict")}<>"unknown"'

arms = sorted({r[0] for r in rows4})
clients = sorted({r[4] for r in rows4})
wsc = wb.create_sheet("Exp4_coverage")
wsc.append(["Client"] + arms)
for cl in clients:
    r = [cl]
    for a in arms:
        dec = (f'COUNTIFS({rng("Exp4", c4, n4, "Client")},"{cl}",{rng("Exp4", c4, n4, "Arm")},"{a}",'
               f'{rng("Exp4", c4, n4, "Decidable")},TRUE)')
        tot = f'COUNTIFS({rng("Exp4", c4, n4, "Client")},"{cl}",{rng("Exp4", c4, n4, "Arm")},"{a}")'
        r.append(f"=IF({tot}=0,\"\",{dec}/{tot})")
    wsc.append(r)
for c in wsc[1]:
    c.font, c.fill = BOLD, HEAD_FILL
for row in wsc.iter_rows(min_row=2):
    for c in row:
        c.font = FONT
        if c.column > 1:
            c.number_format = "0%"
wsc.column_dimensions["A"].width = 36
wsc.freeze_panes = "B2"
wsc["A1"].comment = Comment("Share of connections with a decidable version verdict (allow or violation), "
                            "per client and arm. C9.", "eval")

# -- Exp5 ---------------------------------------------------------------------
order = {"off": 0, "audit": 1, "enforce": 2}
cost = sorted(score.jsonl(R / "exp5/cost/runs.jsonl"), key=lambda r: (order[r["condition"]], r["round"]))
ws5, c5, n5 = sheet("Exp5_cost", ["Condition", "Round", "Connections", "Wall seconds", "Handshakes per s", "Attached",
                                  "Detected", "Dropped", "Blocks"],
                    [[r["condition"], r["round"], r["connections"], r["wall_seconds"], r["cps"], r["attached"],
                      (r["skbouncer"] or {}).get("counters", {}).get("detected"),
                      (r["skbouncer"] or {}).get("counters", {}).get("dropped", 0),
                      (r["skbouncer"] or {}).get("blocks", 0)] for r in cost])
stats = sorted(score.jsonl(R / "exp5/cost-stats/runs.jsonl"), key=lambda r: (order[r["condition"]], r["round"]))
ws5s, c5s, n5s = sheet("Exp5_stats", ["Condition", "Round", "Handshakes per s", "run_cnt", "run_time_ns", "ns per call",
                                      "Dropped"],
                       [[r["condition"], r["round"], r["cps"], (r["skbouncer"] or {}).get("bpf", {}).get("run_cnt") if r["skbouncer"] else None,
                         (r["skbouncer"] or {}).get("bpf", {}).get("run_time_ns") if r["skbouncer"] else None, None,
                         (r["skbouncer"] or {}).get("counters", {}).get("dropped", 0)] for r in stats])
for i in range(2, n5s + 1):
    ws5s[f"F{i}"] = f'=IF(D{i}="","",E{i}/D{i})'
tput = sorted(score.jsonl(R / "exp5/load/throughput.jsonl"), key=lambda r: (order[r["condition"]], r["round"]))
ws5t, c5t, n5t = sheet("Exp5_tput", ["Condition", "Round", "Bytes", "Wall seconds", "Gbps", "Dropped"],
                       [[r["condition"], r["round"], r["bytes"], r["wall_seconds"], r["gbps"],
                         (r["skbouncer"] or {}).get("counters", {}).get("dropped", 0)] for r in tput])
lat = {c: [] for c in order}
latd = 0
for r in score.jsonl(R / "exp5/load/latency.jsonl"):
    lat[r["condition"]].extend(r["ms"])
    latd += (r.get("skbouncer") or {}).get("counters", {}).get("dropped", 0)
wsl = wb.create_sheet("Exp5_latency")
wsl.append(["off (ms)", "audit (ms)", "enforce (ms)"])
for i in range(max(len(v) for v in lat.values())):
    wsl.append([lat[c][i] if i < len(lat[c]) else None for c in order])
for c in wsl[1]:
    c.font, c.fill = BOLD, HEAD_FILL
wsl.freeze_panes = "A2"
nl = max(len(v) for v in lat.values()) + 1
wsl["E1"], wsl["F1"], wsl["G1"], wsl["H1"] = "Statistic", "off", "audit", "enforce"
for j, (name, q) in enumerate([("p50", 0.5), ("p99", 0.99)], 2):
    wsl[f"E{j}"] = name
    for col, src in zip("FGH", "ABC"):
        wsl[f"{col}{j}"] = f"=PERCENTILE({src}2:{src}{nl},{q})"
wsl["E4"], wsl["F4"] = "Dropped events", latd
for col in "ABCEFGH":
    wsl.column_dimensions[col].width = 14
pol = sorted(score.jsonl(R / "exp5/load/policy.jsonl"), key=lambda r: (r["rules"], -r["ciphers"], r["round"]))
ws5p, c5p, n5p = sheet("Exp5_policy", ["Rules", "Ciphers", "Round", "Connections", "Wall seconds", "Handshakes per s",
                                       "Blocks", "Dropped"],
                       [[r["rules"], r["ciphers"], r["round"], r["connections"], r["wall_seconds"], r["cps"],
                         r["skbouncer"]["blocks"], r["skbouncer"]["counters"]["dropped"]] for r in pol])
sat = score.jsonl(R / "exp5/load/saturation.jsonl")
ws5x, c5x, n5x = sheet("Exp5_sat", ["Parallel clients", "Connections", "Wall seconds", "Handshakes per s", "Detected", "Dropped"],
                       [[r["parallel"], r["connections"], r["wall_seconds"], r["cps"], r["counters"]["detected"],
                         r["counters"]["dropped"]] for r in sat])
fp = json.loads((R / "exp5/load/footprint.json").read_text())
ws5f, c5f, n5f = sheet("Exp5_footprint", ["Item", "Value", "Unit"],
                       [["Program, translated", fp["prog"]["bytes_xlated"], "bytes"],
                        ["Program, JIT", fp["prog"]["bytes_jited"], "bytes"],
                        ["Verified instructions", fp["prog"]["verified_insns"], "instructions"]]
                       + [[f"Map {m['name']} ({m['type']})", m["bytes_memlock"], "bytes"] for m in fp["maps"]]
                       + [["Userspace resident memory", fp["rss_kib"], "KiB"],
                          ["Userspace CPU time", fp["cpu_seconds"], "s"], ["Over a run of", fp["seconds"], "s"]],
                       {"Item": 34})

# -- Exp6 ---------------------------------------------------------------------
pols6 = {n: score.load_policy(R / f"exp6/policies/{n}.json") for n in ("min13", "ciphers", "min13_block")}
rows6 = []
for r in score.jsonl(R / "exp6/matrix.jsonl"):
    e = r["events"][0]
    rows6.append([r["label"], r["policy"], e["verdict"],
                  score.reference_verdict(pols6[r["policy"]], e.get("container"), r["reference"]),
                  pols6[r["policy"]]["unknown_block"], r["enforce_received"], None, None, None])
ws6, c6, n6 = sheet("Exp6_matrix", ["Client", "Policy", "skbouncer verdict", "Reference verdict", "on_unknown block",
                                    "Server received", "Blocked", "Should block", "Agrees"], rows6, {"Client": 34})
for i in range(2, n6 + 1):
    ws6[f"G{i}"] = f"=NOT(F{i})"
    ws6[f"H{i}"] = f'=IF(C{i}="unknown",E{i},D{i}="violation")'
    ws6[f"I{i}"] = f"=G{i}=H{i}"
ws6s, c6s, n6s = sheet("Exp6_scenarios", ["Scenario", "Audit verdict", "Enforce: server received", "Blocked"],
                       [[s["name"], s["audit"]["events"][0]["verdict"], s["enforce"]["received"], None]
                        for s in score.jsonl(R / "exp6/scenarios.jsonl")])
for i in range(2, n6s + 1):
    ws6s[f"D{i}"] = f"=NOT(C{i})"
ws6w, c6w, n6w = sheet("Exp6_waits", ["Client", "Seconds until failure", "Capped at 300 s", "Exit status"],
                       [[w["label"], w["seconds"], w["capped"], w["exit"]] for w in score.jsonl(R / "exp6/waits.jsonl")],
                       {"Client": 34})

# -- Exp7 ---------------------------------------------------------------------
rows7 = []
for n in ("n10", "n50"):
    for c in score.jsonl(R / f"exp7/{n}/containers.jsonl"):
        ev = c["events"]
        rows7.append([n[1:], c["name"], c["shared"], c["veth"], c["ip"], c["dport"], c["handshakes_ok"], len(ev),
                      sum(e.get("container") == c["name"] for e in ev), sum(e.get("container") is None for e in ev)])
ws7, c7, n7 = sheet("Exp7", ["Containers in run", "Container", "Shares a namespace", "veth", "IP", "Server port",
                             "Handshakes ok", "Events", "Events naming it", "Unnamed events"], rows7, {"Container": 28})

# -- Environment and About ----------------------------------------------------
env = [l.split(": ", 1) for l in (R / "env-eval-1.1.txt").read_text().splitlines() if ": " in l]
sheet("Environment", ["Item", "Value"], env, {"Item": 34, "Value": 90})

# -- Criteria -----------------------------------------------------------------
X3 = lambda h: rng("Exp3", c3c, n3, h)
X4 = lambda h: rng("Exp4", c4, n4, h)
med = lambda sh, col, a, b: f"MEDIAN({sh}!${col}${a}:${col}${b})"
cost_rows = {c: [i + 2 for i, r in enumerate(cost) if r["condition"] == c] for c in order}
tput_rows = {c: [i + 2 for i, r in enumerate(tput) if r["condition"] == c] for c in order}
pol_rows = {}
for i, r in enumerate(pol):
    pol_rows.setdefault((r["rules"], r["ciphers"]), []).append(i + 2)
cm = {c: med("Exp5_cost", "E", v[0], v[-1]) for c, v in cost_rows.items()}
tm = {c: med("Exp5_tput", "E", v[0], v[-1]) for c, v in tput_rows.items()}
pm = {k: med("Exp5_policy", "F", v[0], v[-1]) for k, v in pol_rows.items()}
sc = score.score_exp4(R)
C = [
    ("C1", "Functional correctness", "14 of 14", f'=COUNTIF(Exp1!$D$2:$D${n1},"PASS")', 14, "eq"),
    ("C2", "Detection rate", "100%", f'=COUNTIF({X3("Events")},">=1")', f'=COUNTA({X3("Run")})', "eq"),
    ("C3", "False positives", "0", "=Exp1!G6", 0, "eq"),
    ("C4", "Cipher list accuracy", "100%", f'=COUNTIF({X3("Cipher match")},TRUE)', f'=COUNTIF({X3("ciphers_seen")},TRUE)', "eq"),
    ("C5", "Version accuracy", "100%", f'=COUNTIF({X3("Version match")},TRUE)', f'=COUNTIF({X3("versions_seen")},TRUE)', "eq"),
    ("C6", "False certainty", "0", f'=COUNTIF({X3("Cipher match")},FALSE)+COUNTIF({X3("Version match")},FALSE)'
     f'+COUNTIF({X3("record_len match")},FALSE)+COUNTIF({X3("legacy match")},FALSE)', 0, "eq"),
    ("C7", "Container attribution", "100%", f'=COUNTIF({X3("Named correctly")},TRUE)', f'=COUNTA({X3("Run")})', "eq"),
    ("C8", "Truncation reporting (split flagged)", "100%", f'=COUNTIFS({X4("Split (truth)")},TRUE,{X4("Flagged split")},TRUE)',
     f'=COUNTIF({X4("Split (truth)")},TRUE)', "eq"),
    ("C8", "Truncation reporting (fits but flagged)", "0", f'=COUNTIFS({X4("Split (truth)")},FALSE,{X4("Flagged split")},TRUE)', 0, "eq"),
    ("C9", "Fragmentation coverage", "tested against H1 to H3",
     "H1 holds except " + "; ".join(sc["h1_exceptions"]) if sc["h1_exceptions"] else "H1 holds", "see Exp4_coverage", "rep"),
    ("C10", "Verdict agreement", "100%", f'=COUNTIF({X3("Verdict agrees")},TRUE)', f'=COUNTIF({X3("Verdict")},"<>unknown")', "eq"),
    ("C11", "Enforcement accuracy (rows 8 to 10)", "all pass", '=COUNTIF(Exp1!$D$9:$D$11,"PASS")', 3, "eq"),
    ("C12", "Safe failure (rows 9, 11 to 13)", "all hold",
     '=COUNTIF(Exp1!$D$10,"PASS")+COUNTIF(Exp1!$D$12:$D$14,"PASS")', 4, "eq"),
    ("C13", "Handshake overhead, worst change", "under 5%",
     f"=MAX(ABS({cm['audit']}/{cm['off']}-1),ABS({cm['enforce']}/{cm['off']}-1))", 0.05, "lt"),
    ("C14", "Kernel cost per call, enforce (ns)", "reported",
     '=AVERAGEIFS(Exp5_stats!$F$2:$F$%d,Exp5_stats!$A$2:$A$%d,"enforce")' % (n5s, n5s), "", "rep"),
    ("C15", "Event loss in experiment 5", "0",
     f"=SUM(Exp5_cost!$H$2:$H${n5})+SUM(Exp5_stats!$G$2:$G${n5s})+SUM(Exp5_tput!$F$2:$F${n5t})"
     f"+Exp5_latency!F4+SUM(Exp5_policy!$H$2:$H${n5p})", 0, "eq"),
    ("C16", "Robustness: false complete fields", "0", f"=SUM(Exp3_fuzz!$I$2:$I${nf})", 0, "eq"),
    ("C17", "Data-path overhead, worst change", "under 5%",
     f"=MAX(ABS({tm['audit']}/{tm['off']}-1),ABS({tm['enforce']}/{tm['off']}-1))", 0.05, "lt"),
    ("C18", "Tail latency p99, worst change", "under 10%", "=MAX(Exp5_latency!G3/Exp5_latency!F3-1,Exp5_latency!H3/Exp5_latency!F3-1)",
     0.10, "lt"),
    ("C19", "Policy scalability, 1 to 64 rules", "under 5%",
     f"=MAX(ABS({pm[(16, 30)]}/{pm[(1, 30)]}-1),ABS({pm[(64, 30)]}/{pm[(1, 30)]}-1))", 0.05, "lt"),
    ("C19", "Policy scalability, 1 against 30 ciphers", "under 5%", f"=ABS({pm[(1, 1)]}/{pm[(1, 30)]}-1)", 0.05, "lt"),
    ("C20", "Saturation: highest rate with no drop (per s)", "reported",
     f"=_xlfn.MAXIFS(Exp5_sat!$D$2:$D${n5x},Exp5_sat!$F$2:$F${n5x},0)", "", "rep"),
    ("C21", "Footprint: JIT bytes", "reported", "=Exp5_footprint!B3", "", "rep"),
    ("C22", "Enforcement at scale", "100%", f"=COUNTIF(Exp6_matrix!$I$2:$I${n6},TRUE)", f"=COUNTA(Exp6_matrix!$A$2:$A${n6})", "eq"),
    ("C23", "Scenario validation", "all", f'=COUNTIFS(Exp6_scenarios!$B$2:$B${n6s},"<>allow",Exp6_scenarios!$D$2:$D${n6s},TRUE)',
     f"=COUNTA(Exp6_scenarios!$A$2:$A${n6s})", "eq"),
    ("C24", "Attribution under concurrency", "100%, none unnamed", f"=SUM(Exp7!$I$2:$I${n7})", f"=SUM(Exp7!$H$2:$H${n7})", "eq"),
    ("C25", "Grade agreement", "100%", f'=COUNTIF({X3("Grade agrees")},TRUE)', f'=COUNTIF({X3("Read in full")},TRUE)', "eq"),
]
wsC = wb.create_sheet("Criteria", 0)
wsC.append(["ID", "Criterion", "Target", "Measured", "Out of / threshold", "Status"])
for i, (cid, name, target, measured, of, kind) in enumerate(C, 2):
    wsC.append([cid, name, target, measured, of, None])
    if kind == "eq":
        wsC[f"F{i}"] = f'=IF(D{i}=E{i},"met","not met")'
    elif kind == "lt":
        wsC[f"F{i}"] = f'=IF(D{i}<E{i},"met","not met")'
        wsC[f"D{i}"].number_format = wsC[f"E{i}"].number_format = "0.00%"
    else:
        wsC[f"F{i}"] = "reported"
for c in wsC[1]:
    c.font, c.fill = BOLD, HEAD_FILL
for row in wsC.iter_rows(min_row=2):
    for c in row:
        c.font = FONT
for col, w in zip("ABCDEF", (6, 44, 22, 60, 20, 10)):
    wsC.column_dimensions[col].width = w
wsC.freeze_panes = "A2"
wsC["D11"].comment = Comment("H2 and H3 hold. Per client and arm shares in Exp4_coverage.", "eval")

about = [
    ["Source", "eval/results/results-eval-1.1.tar.xz (SHA-256 02a946ed...7a27), tag eval-1.1"],
    ["Run", "2026-10-03 on EC2 c7i.2xlarge, ap-southeast-1, kernel 7.0.0-1013-aws (sheet Environment)"],
    ["Criteria", "eval/PLAN.md. Figures on the Criteria sheet are formulas over the data sheets."],
    ["Reference verdicts", "eval/score.py reference_verdict(): policy/README.md applied to the server's own decode"],
    ["Grades", "Durumeric et al., NDSS 2017, section V-A; cipher properties from IANA's registry"],
    ["Fuzz decode", "tests/check_m7_3.py decode(): an independent decode of each mutated ClientHello as sent"],
    ["Experiment 5 raw events", "Per-handshake event streams are in the archive only; the sheets hold per-run counters"],
]
sheet("About", ["Item", "Note"], about, {"Item": 24, "Note": 100})
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT)
